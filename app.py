import io
import re
import tempfile
from decimal import Decimal, InvalidOperation, ROUND_CEILING
from pathlib import Path

import pandas as pd
import pymupdf
import streamlit as st
from openpyxl import Workbook, load_workbook
from openpyxl.drawing.image import Image as ExcelImage
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter


# ==========================================================
# CONFIGURAÇÃO
# ==========================================================

st.set_page_config(
    page_title="Atualizador de Preços",
    page_icon="💰",
    layout="wide",
)

st.title("💰 Atualizador de Preços")
st.caption(
    "PDF → PDF, PDF → Excel e Excel → Excel."
)


# ==========================================================
# PREÇOS E ARREDONDAMENTO
# ==========================================================

PADRAO_PRECO = re.compile(
    r"R\$\s*\d{1,3}(?:\.\d{3})*,\d{2}"
)


def arredondar_sempre_para_cima(valor):
    """
    Exemplos:
    10,01 -> 11
    10,99 -> 11
    20,00 -> 20
    """
    return int(
        Decimal(str(valor)).quantize(
            Decimal("1"),
            rounding=ROUND_CEILING,
        )
    )


def texto_preco_para_decimal(texto_preco):
    """
    Converte:
    R$ 926,82 -> Decimal("926.82")
    R$ 1.073,88 -> Decimal("1073.88")
    """
    texto = str(texto_preco).upper()
    texto = texto.replace("R$", "").strip()
    texto = texto.replace(".", "").replace(",", ".")

    try:
        return Decimal(texto)
    except InvalidOperation:
        return None


def formatar_preco_brasileiro(valor_inteiro):
    """
    Converte:
    2210 -> R$ 2.210
    """
    numero = f"{valor_inteiro:,}"
    numero = numero.replace(",", ".")
    return f"R$ {numero}"


def calcular_novo_preco(texto_preco, multiplicador):
    """
    Novo preço = teto(preço original x multiplicador)
    """
    valor_original = texto_preco_para_decimal(texto_preco)

    if valor_original is None:
        return None

    resultado = valor_original * Decimal(str(multiplicador))

    return arredondar_sempre_para_cima(resultado)


# ==========================================================
# PDF -> PDF
# ==========================================================

def rgb_inteiro_para_tupla(cor):
    if not isinstance(cor, int):
        return (0, 0, 0)

    vermelho = ((cor >> 16) & 255) / 255
    verde = ((cor >> 8) & 255) / 255
    azul = (cor & 255) / 255

    return (vermelho, verde, azul)


def escolher_fonte_pdf(nome_fonte_original):
    nome = str(nome_fonte_original).lower()

    if "bold" in nome:
        return "hebo"

    if "times" in nome:
        return "tiro"

    if "courier" in nome:
        return "cour"

    return "helv"


def medir_largura_texto(texto, fonte, tamanho):
    try:
        return pymupdf.get_text_length(
            texto,
            fontname=fonte,
            fontsize=tamanho,
        )
    except Exception:
        return len(texto) * tamanho * 0.55


def localizar_precos_na_pagina(pagina):
    encontrados = []

    estrutura = pagina.get_text("dict")

    for bloco in estrutura.get("blocks", []):
        if bloco.get("type") != 0:
            continue

        for linha in bloco.get("lines", []):
            for span in linha.get("spans", []):
                texto_span = span.get("text", "")

                if not texto_span:
                    continue

                correspondencias = list(
                    PADRAO_PRECO.finditer(texto_span)
                )

                if not correspondencias:
                    continue

                x0, y0, x1, y1 = span["bbox"]
                largura_span = x1 - x0
                tamanho_span = max(len(texto_span), 1)

                for correspondencia in correspondencias:
                    inicio = correspondencia.start()
                    fim = correspondencia.end()

                    if inicio == 0 and fim == len(texto_span):
                        caixa_preco = pymupdf.Rect(
                            x0,
                            y0,
                            x1,
                            y1,
                        )
                    else:
                        x_inicio = x0 + (
                            largura_span * inicio / tamanho_span
                        )

                        x_fim = x0 + (
                            largura_span * fim / tamanho_span
                        )

                        caixa_preco = pymupdf.Rect(
                            x_inicio,
                            y0,
                            x_fim,
                            y1,
                        )

                    encontrados.append(
                        {
                            "texto_original": correspondencia.group(0),
                            "rect": caixa_preco,
                            "fonte_original": span.get("font", "helv"),
                            "tamanho_fonte": span.get("size", 10),
                            "cor": rgb_inteiro_para_tupla(
                                span.get("color", 0)
                            ),
                        }
                    )

    return encontrados


def colocar_fundo_branco(pagina):
    pagina.draw_rect(
        pagina.rect,
        color=None,
        fill=(1, 1, 1),
        overlay=False,
    )


def atualizar_pdf_mantendo_layout(
    arquivo_pdf,
    multiplicador,
    deixar_fundo_branco=True,
):
    documento = pymupdf.open(
        stream=arquivo_pdf,
        filetype="pdf",
    )

    conferencias = []

    if deixar_fundo_branco:
        for pagina in documento:
            colocar_fundo_branco(pagina)

    for numero_pagina, pagina in enumerate(documento, start=1):
        precos_pagina = localizar_precos_na_pagina(pagina)
        substituicoes = []

        for preco in precos_pagina:
            texto_original = preco["texto_original"]

            novo_valor = calcular_novo_preco(
                texto_original,
                multiplicador,
            )

            if novo_valor is None:
                continue

            texto_novo = formatar_preco_brasileiro(
                novo_valor
            )

            caixa_original = preco["rect"]

            caixa_redacao = pymupdf.Rect(
                caixa_original.x0 - 1.5,
                caixa_original.y0 - 1.0,
                caixa_original.x1 + 1.5,
                caixa_original.y1 + 1.0,
            )

            pagina.add_redact_annot(
                caixa_redacao,
                fill=(1, 1, 1),
            )

            substituicoes.append(
                {
                    "texto_original": texto_original,
                    "texto_novo": texto_novo,
                    "caixa_original": caixa_original,
                    "fonte": escolher_fonte_pdf(
                        preco["fonte_original"]
                    ),
                    "tamanho": preco["tamanho_fonte"],
                    "cor": preco["cor"],
                }
            )

        if substituicoes:
            pagina.apply_redactions(
                images=pymupdf.PDF_REDACT_IMAGE_NONE,
                graphics=pymupdf.PDF_REDACT_LINE_ART_NONE,
                text=pymupdf.PDF_REDACT_TEXT_REMOVE,
            )

        for item in substituicoes:
            caixa = item["caixa_original"]
            texto_novo = item["texto_novo"]
            fonte = item["fonte"]
            cor = item["cor"]
            tamanho_ajustado = item["tamanho"]

            largura_original = caixa.width
            largura_nova = medir_largura_texto(
                texto_novo,
                fonte,
                tamanho_ajustado,
            )

            while (
                largura_nova > largura_original
                and tamanho_ajustado > 5
            ):
                tamanho_ajustado -= 0.2

                largura_nova = medir_largura_texto(
                    texto_novo,
                    fonte,
                    tamanho_ajustado,
                )

            caixa_texto = pymupdf.Rect(
                caixa.x0,
                caixa.y0 - 0.5,
                max(
                    caixa.x1,
                    caixa.x0 + largura_nova + 3,
                ),
                caixa.y1 + 2,
            )

            resultado_texto = pagina.insert_textbox(
                caixa_texto,
                texto_novo,
                fontname=fonte,
                fontsize=tamanho_ajustado,
                color=cor,
                align=pymupdf.TEXT_ALIGN_LEFT,
                overlay=True,
            )

            if resultado_texto < 0:
                pagina.insert_text(
                    pymupdf.Point(
                        caixa.x0,
                        caixa.y1 - 1,
                    ),
                    texto_novo,
                    fontname=fonte,
                    fontsize=tamanho_ajustado,
                    color=cor,
                    overlay=True,
                )

            conferencias.append(
                {
                    "Página": numero_pagina,
                    "Preço original": item["texto_original"],
                    "Preço atualizado": item["texto_novo"],
                }
            )

    arquivo_final = io.BytesIO()

    documento.save(
        arquivo_final,
        garbage=4,
        deflate=True,
    )

    documento.close()
    arquivo_final.seek(0)

    return arquivo_final.getvalue(), conferencias


# ==========================================================
# EXCEL -> EXCEL
# ==========================================================

def converter_valor_excel_para_decimal(valor):
    """
    Aceita:
    926.82
    926,82
    1.073,88
    R$ 1.073,88
    """
    if valor is None:
        return None

    if isinstance(valor, (int, float, Decimal)):
        return Decimal(str(valor))

    texto = str(valor).strip()

    if not texto:
        return None

    texto = texto.upper().replace("R$", "").strip()
    texto = texto.replace(" ", "")

    if "," in texto and "." in texto:
        texto = texto.replace(".", "").replace(",", ".")
    elif "," in texto:
        texto = texto.replace(",", ".")

    try:
        return Decimal(texto)
    except InvalidOperation:
        return None


def parece_preco_excel(valor):
    """
    Aceita números e preços.
    Evita alterar medidas e códigos.
    """
    if isinstance(valor, (int, float, Decimal)):
        return True

    if not isinstance(valor, str):
        return False

    padrao = (
        r"^\s*"
        r"(R\$\s*)?"
        r"\d{1,3}(?:\.\d{3})*"
        r"(?:,\d{1,2})?"
        r"\s*$"
    )

    return bool(re.match(padrao, valor))


def letra_para_numero_coluna(letra_coluna):
    """
    E -> 5
    AA -> 27
    """
    numero_coluna = 0

    for caractere in letra_coluna.upper():
        numero_coluna = (
            numero_coluna * 26
            + ord(caractere)
            - ord("A")
            + 1
        )

    return numero_coluna


def calcular_formula_simples(formula, planilha):
    """
    Calcula fórmulas simples de Excel.

    Aceita:
    +, -, *, /, parênteses
    referências: A1, E5, AA22

    Exemplos:
    =E5+100
    =E5*1.10
    =(E5+100)*1.05

    Não aceita:
    =SOMA(E5:E10)
    =SE(A1>0;1;0)
    =PROCV(...)
    """
    expressao = formula[1:].strip()

    padrao_celula = re.compile(
        r"\$?([A-Z]{1,3})\$?(\d+)"
    )

    def substituir_referencia(match):
        letra = match.group(1)
        numero_linha = int(match.group(2))

        numero_coluna = letra_para_numero_coluna(
            letra
        )

        valor_referencia = planilha.cell(
            numero_linha,
            numero_coluna,
        ).value

        if valor_referencia is None:
            return "0"

        if (
            isinstance(valor_referencia, str)
            and valor_referencia.startswith("=")
        ):
            raise ValueError(
                "A fórmula depende de outra fórmula: "
                + letra
                + str(numero_linha)
            )

        valor_decimal = converter_valor_excel_para_decimal(
            valor_referencia
        )

        if valor_decimal is None:
            raise ValueError(
                "A célula "
                + letra
                + str(numero_linha)
                + " não possui valor numérico."
            )

        return str(valor_decimal)

    expressao = padrao_celula.sub(
        substituir_referencia,
        expressao,
    )

    expressao = expressao.replace(",", ".")

    if not re.fullmatch(
        r"[0-9\.\+\-\*\/\(\)\s]+",
        expressao,
    ):
        raise ValueError(
            "Fórmula complexa não suportada."
        )

    try:
        resultado = eval(
            expressao,
            {"__builtins__": {}},
            {},
        )

        return Decimal(str(resultado))

    except Exception as erro:
        raise ValueError(
            "Erro ao calcular fórmula: " + str(erro)
        )


def calcular_formula_com_multiplicador(
    formula,
    planilha,
    multiplicador,
):
    """
    Opção 2:
    calcula toda a fórmula, multiplica o resultado e arredonda para cima.

    Exemplo:
    =E5+100

    Resultado:
    teto((resultado de E5+100) x multiplicador)
    """
    resultado_formula = calcular_formula_simples(
        formula,
        planilha,
    )

    resultado_multiplicado = (
        resultado_formula
        * Decimal(str(multiplicador))
    )

    return arredondar_sempre_para_cima(
        resultado_multiplicado
    )


def atualizar_excel(
    arquivo_excel,
    nome_aba,
    linha_inicio,
    numero_coluna,
    multiplicador,
):
    """
    Atualiza somente uma coluna escolhida pela letra.

    Para valores diretos:
    teto(valor x multiplicador)

    Para fórmulas simples:
    calcula o resultado, multiplica, arredonda para cima
    e grava o número final na célula.
    """
    workbook = load_workbook(
        io.BytesIO(arquivo_excel),
        data_only=False,
        keep_links=True,
        rich_text=True,
    )

    planilha = workbook[nome_aba]
    conferencias = []

    for numero_linha in range(
        linha_inicio,
        planilha.max_row + 1,
    ):
        celula = planilha.cell(
            numero_linha,
            numero_coluna,
        )

        valor_original = celula.value

        if valor_original is None:
            continue

        # --------------------------------------------------
        # FÓRMULAS SIMPLES
        # --------------------------------------------------
        if (
            isinstance(valor_original, str)
            and valor_original.startswith("=")
        ):
            try:
                resultado_final = calcular_formula_com_multiplicador(
                    valor_original,
                    planilha,
                    multiplicador,
                )

                # Remove fórmula e deixa somente número final.
                celula.value = resultado_final

                formato_anterior = str(celula.number_format)

                if (
                    "R$" in formato_anterior
                    or "$" in formato_anterior
                ):
                    celula.number_format = "R$ #,##0"

                conferencias.append(
                    {
                        "Linha": numero_linha,
                        "Célula": celula.coordinate,
                        "Valor original": valor_original,
                        "Valor atualizado": resultado_final,
                        "Tipo": "Fórmula calculada",
                    }
                )

            except Exception as erro_formula:
                conferencias.append(
                    {
                        "Linha": numero_linha,
                        "Célula": celula.coordinate,
                        "Valor original": valor_original,
                        "Valor atualizado": "Fórmula mantida",
                        "Tipo": "Não suportada: "
                        + str(erro_formula),
                    }
                )

            continue

        # --------------------------------------------------
        # NÚMEROS OU PREÇOS DIRETOS
        # --------------------------------------------------
        if not parece_preco_excel(valor_original):
            continue

        valor_decimal = converter_valor_excel_para_decimal(
            valor_original
        )

        if valor_decimal is None:
            continue

        novo_valor = arredondar_sempre_para_cima(
            valor_decimal * Decimal(str(multiplicador))
        )

        if (
            isinstance(valor_original, str)
            and "R$" in valor_original.upper()
        ):
            celula.value = formatar_preco_brasileiro(
                novo_valor
            )
        else:
            celula.value = novo_valor

            formato_anterior = str(celula.number_format)

            if (
                "R$" in formato_anterior
                or "$" in formato_anterior
            ):
                celula.number_format = "R$ #,##0"

        conferencias.append(
            {
                "Linha": numero_linha,
                "Célula": celula.coordinate,
                "Valor original": valor_original,
                "Valor atualizado": celula.value,
                "Tipo": "Valor multiplicado",
            }
        )

    arquivo_final = io.BytesIO()
    workbook.save(arquivo_final)
    arquivo_final.seek(0)

    return arquivo_final.getvalue(), conferencias


# ==========================================================
# PDF -> EXCEL
# ==========================================================

def criar_cabecalho(planilha, titulos):
    for numero_coluna, titulo in enumerate(titulos, start=1):
        celula = planilha.cell(
            1,
            numero_coluna,
            titulo,
        )

        celula.font = Font(
            bold=True,
            color="FFFFFF",
        )

        celula.fill = PatternFill(
            fill_type="solid",
            fgColor="1F4E78",
        )

        celula.alignment = Alignment(
            horizontal="center",
            vertical="center",
            wrap_text=True,
        )


def inserir_pagina_pdf_como_imagem(planilha, pagina):
    pix = pagina.get_pixmap(
        matrix=pymupdf.Matrix(1.4, 1.4),
        alpha=False,
    )

    imagem_bytes = pix.tobytes("png")

    arquivo_temporario = tempfile.NamedTemporaryFile(
        suffix=".png",
        delete=False,
    )

    arquivo_temporario.write(imagem_bytes)
    arquivo_temporario.close()

    imagem_excel = ExcelImage(
        arquivo_temporario.name
    )

    largura_maxima = 1100

    if imagem_excel.width > largura_maxima:
        proporcao = largura_maxima / imagem_excel.width

        imagem_excel.width = int(
            imagem_excel.width * proporcao
        )

        imagem_excel.height = int(
            imagem_excel.height * proporcao
        )

    planilha.add_image(imagem_excel, "A1")
    planilha.column_dimensions["A"].width = 20


def adicionar_tabela_extraida(
    planilha,
    tabela,
    linha_inicial,
    numero_pagina,
):
    dados = tabela.extract()

    if not dados:
        return linha_inicial

    celula_titulo = planilha.cell(
        linha_inicial,
        1,
        "Tabela encontrada - Página "
        + str(numero_pagina),
    )

    celula_titulo.font = Font(
        bold=True,
        color="FFFFFF",
    )

    celula_titulo.fill = PatternFill(
        fill_type="solid",
        fgColor="4F81BD",
    )

    linha_inicial += 1

    for linha_pdf in dados:
        for numero_coluna, valor in enumerate(
            linha_pdf,
            start=1,
        ):
            celula = planilha.cell(
                linha_inicial,
                numero_coluna,
                valor,
            )

            celula.alignment = Alignment(
                vertical="top",
                wrap_text=True,
            )

        linha_inicial += 1

    return linha_inicial + 2


def converter_pdf_para_excel(arquivo_pdf):
    documento = pymupdf.open(
        stream=arquivo_pdf,
        filetype="pdf",
    )

    workbook = Workbook()

    planilha_tabelas = workbook.active
    planilha_tabelas.title = "Tabelas extraídas"

    planilha_texto = workbook.create_sheet(
        "Texto por página"
    )

    criar_cabecalho(
        planilha_texto,
        [
            "Página",
            "Texto extraído",
        ],
    )

    planilha_texto.column_dimensions["A"].width = 12
    planilha_texto.column_dimensions["B"].width = 100
    planilha_texto.freeze_panes = "A2"

    linha_texto = 2
    linha_tabelas = 1
    quantidade_tabelas = 0

    for numero_pagina, pagina in enumerate(
        documento,
        start=1,
    ):
        texto_pagina = pagina.get_text("text")

        for linha in texto_pagina.splitlines():
            linha = linha.strip()

            if not linha:
                continue

            planilha_texto.cell(
                linha_texto,
                1,
                numero_pagina,
            )

            celula_texto = planilha_texto.cell(
                linha_texto,
                2,
                linha,
            )

            celula_texto.alignment = Alignment(
                vertical="top",
                wrap_text=True,
            )

            linha_texto += 1

        try:
            resultado_tabelas = pagina.find_tables()

            for tabela in resultado_tabelas.tables:
                linha_tabelas = adicionar_tabela_extraida(
                    planilha_tabelas,
                    tabela,
                    linha_tabelas,
                    numero_pagina,
                )

                quantidade_tabelas += 1

        except Exception:
            pass

        planilha_visual = workbook.create_sheet(
            "Visual página " + str(numero_pagina)
        )

        inserir_pagina_pdf_como_imagem(
            planilha_visual,
            pagina,
        )

    planilha_tabelas.column_dimensions["A"].width = 32

    for numero_coluna in range(2, 20):
        letra = get_column_letter(numero_coluna)
        planilha_tabelas.column_dimensions[letra].width = 22

    planilha_info = workbook.create_sheet("Informações")

    planilha_info["A1"] = "Conversão de PDF para Excel"
    planilha_info["A1"].font = Font(
        bold=True,
        size=14,
    )

    planilha_info["A3"] = "Tabelas identificadas"
    planilha_info["B3"] = quantidade_tabelas

    planilha_info["A5"] = "Observação"
    planilha_info["A5"].font = Font(bold=True)

    planilha_info["B5"] = (
        "A aba 'Tabelas extraídas' contém tabelas reconhecidas. "
        "A aba 'Texto por página' contém conteúdo editável. "
        "As abas 'Visual página N' preservam a aparência das páginas "
        "originais como imagem."
    )

    planilha_info["B5"].alignment = Alignment(
        wrap_text=True,
        vertical="top",
    )

    planilha_info.column_dimensions["A"].width = 28
    planilha_info.column_dimensions["B"].width = 100
    planilha_info.row_dimensions[5].height = 75

    arquivo_final = io.BytesIO()
    workbook.save(arquivo_final)
    arquivo_final.seek(0)

    documento.close()

    return arquivo_final.getvalue(), quantidade_tabelas


# ==========================================================
# INTERFACE PRINCIPAL
# ==========================================================

arquivo = st.file_uploader(
    "Envie um arquivo PDF ou Excel",
    type=["pdf", "xlsx"],
)

multiplicador = st.number_input(
    "Multiplicador de preço",
    min_value=0.000001,
    value=2.383949988,
    step=0.000001,
    format="%.9f",
)

if arquivo is None:
    st.info("Envie um arquivo para começar.")
    st.stop()

nome_arquivo = arquivo.name
extensao = Path(nome_arquivo).suffix.lower()
conteudo_arquivo = arquivo.getvalue()


# ==========================================================
# TELA PDF
# ==========================================================

if extensao == ".pdf":
    st.subheader("PDF")

    modo_pdf = st.radio(
        "O que deseja gerar?",
        [
            "PDF com preços atualizados",
            "Excel convertido do PDF",
        ],
    )

    if modo_pdf == "PDF com preços atualizados":
        deixar_fundo_branco = st.checkbox(
            "Adicionar fundo branco atrás dos elementos",
            value=True,
        )

        if st.button(
            "Gerar PDF com preços atualizados",
            type="primary",
        ):
            try:
                with st.spinner(
                    "Atualizando preços no PDF..."
                ):
                    resultado, conferencias = (
                        atualizar_pdf_mantendo_layout(
                            arquivo_pdf=conteudo_arquivo,
                            multiplicador=multiplicador,
                            deixar_fundo_branco=deixar_fundo_branco,
                        )
                    )

                if not conferencias:
                    st.warning(
                        "Nenhum preço foi identificado. "
                        "O PDF pode ser escaneado ou usar outro padrão."
                    )
                    st.stop()

                st.success(
                    "PDF pronto. "
                    + str(len(conferencias))
                    + " preço(s) atualizado(s)."
                )

                st.dataframe(
                    pd.DataFrame(conferencias),
                    use_container_width=True,
                    hide_index=True,
                )

                nome_saida = (
                    Path(nome_arquivo).stem
                    + "_precos_atualizados.pdf"
                )

                st.download_button(
                    label="⬇️ Baixar PDF atualizado",
                    data=resultado,
                    file_name=nome_saida,
                    mime="application/pdf",
                )

            except Exception as erro:
                st.error(
                    "Não foi possível processar o PDF. "
                    "Detalhe técnico: " + str(erro)
                )

    elif modo_pdf == "Excel convertido do PDF":
        if st.button(
            "Converter PDF para Excel",
            type="primary",
        ):
            try:
                with st.spinner(
                    "Extraindo textos, tabelas e páginas..."
                ):
                    resultado, quantidade_tabelas = (
                        converter_pdf_para_excel(
                            conteudo_arquivo
                        )
                    )

                st.success(
                    "Excel pronto. "
                    + str(quantidade_tabelas)
                    + " tabela(s) identificada(s)."
                )

                nome_saida = (
                    Path(nome_arquivo).stem
                    + "_convertido.xlsx"
                )

                st.download_button(
                    label="⬇️ Baixar Excel convertido",
                    data=resultado,
                    file_name=nome_saida,
                    mime=(
                        "application/vnd.openxmlformats-officedocument."
                        "spreadsheetml.sheet"
                    ),
                )

            except Exception as erro:
                st.error(
                    "Não foi possível converter o PDF para Excel. "
                    "Detalhe técnico: " + str(erro)
                )


# ==========================================================
# TELA EXCEL — SEM TÍTULOS
# ==========================================================

elif extensao == ".xlsx":
    st.subheader("Excel → Excel")

    st.info(
        "Não é necessário haver título de coluna. "
        "Escolha a aba, a primeira linha de valores e a letra da coluna."
    )

    try:
        workbook_preview = load_workbook(
            io.BytesIO(conteudo_arquivo),
            data_only=False,
            keep_links=True,
            rich_text=True,
        )

        nome_aba = st.selectbox(
            "1. Escolha a aba",
            workbook_preview.sheetnames,
        )

        planilha_preview = workbook_preview[nome_aba]

        linha_inicio = st.number_input(
            "2. Linha onde começam os preços",
            min_value=1,
            max_value=max(1, planilha_preview.max_row),
            value=1,
            step=1,
            help=(
                "Exemplo: se o primeiro preço estiver em E5, informe 5."
            ),
        )

        letras_colunas = [
            get_column_letter(numero_coluna)
            for numero_coluna in range(
                1,
                planilha_preview.max_column + 1,
            )
        ]

        letra_coluna = st.selectbox(
            "3. Escolha a letra da coluna dos preços",
            letras_colunas,
        )

        numero_coluna = letra_para_numero_coluna(
            letra_coluna
        )

        st.caption(
            "Coluna selecionada: "
            + letra_coluna
            + ". Células a partir da linha "
            + str(linha_inicio)
            + " serão analisadas."
        )

        dados_previa = []

        for numero_linha in range(
            linha_inicio,
            min(
                planilha_preview.max_row + 1,
                linha_inicio + 10,
            ),
        ):
            valor_original = planilha_preview.cell(
                numero_linha,
                numero_coluna,
            ).value

            if valor_original is None:
                continue

            if (
                isinstance(valor_original, str)
                and valor_original.startswith("=")
            ):
                try:
                    novo_valor = calcular_formula_com_multiplicador(
                        valor_original,
                        planilha_preview,
                        multiplicador,
                    )
                except Exception as erro_formula:
                    novo_valor = (
                        "Fórmula não calculada: "
                        + str(erro_formula)
                    )
            else:
                valor_decimal = converter_valor_excel_para_decimal(
                    valor_original
                )

                if valor_decimal is None:
                    novo_valor = "Não será alterado"
                else:
                    novo_valor = arredondar_sempre_para_cima(
                        valor_decimal * Decimal(str(multiplicador))
                    )

            dados_previa.append(
                {
                    "Linha": numero_linha,
                    "Célula": letra_coluna + str(numero_linha),
                    "Valor original": valor_original,
                    "Novo valor": novo_valor,
                }
            )

        if dados_previa:
            st.subheader("Prévia")
            st.dataframe(
                pd.DataFrame(dados_previa),
                use_container_width=True,
                hide_index=True,
            )
        else:
            st.warning(
                "Nenhum valor foi encontrado nas primeiras linhas "
                "dessa coluna. Verifique a coluna e a linha inicial."
            )

        if st.button(
            "Gerar Excel atualizado",
            type="primary",
        ):
            with st.spinner(
                "Atualizando coluna " + letra_coluna + "..."
            ):
                resultado, conferencias = atualizar_excel(
                    arquivo_excel=conteudo_arquivo,
                    nome_aba=nome_aba,
                    linha_inicio=linha_inicio,
                    numero_coluna=numero_coluna,
                    multiplicador=multiplicador,
                )

            if not conferencias:
                st.warning(
                    "Nenhuma célula foi alterada. Verifique a seleção."
                )
                st.stop()

            st.success(
                "Excel pronto. "
                + str(len(conferencias))
                + " célula(s) processada(s)."
            )

            st.subheader("Conferência")
            st.dataframe(
                pd.DataFrame(conferencias),
                use_container_width=True,
                hide_index=True,
            )

            nome_saida = (
                Path(nome_arquivo).stem
                + "_precos_atualizados.xlsx"
            )

            st.download_button(
                label="⬇️ Baixar Excel atualizado",
                data=resultado,
                file_name=nome_saida,
                mime=(
                    "application/vnd.openxmlformats-officedocument."
                    "spreadsheetml.sheet"
                ),
            )

    except Exception as erro:
        st.error(
            "Não foi possível abrir ou salvar este Excel. "
            "Detalhe técnico: " + str(erro)
        )
